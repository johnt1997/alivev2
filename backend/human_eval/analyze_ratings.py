"""
Wertet die ausgefuellten Bewertungsboegen aus.

  1. Uebereinstimmung Mensch-Mensch: quadratisch gewichtetes Cohen's kappa
  2. Uebereinstimmung Mensch-Judge:  Spearman rho zwischen dem Mittel beider
     Rater und der RAGAS-Faithfulness (gesamt und je Quelle)

Aufruf (aus backend/):  python human_eval/analyze_ratings.py
Erwartet Rater_A.xlsx, Rater_B.xlsx und key_DO_NOT_SHARE.csv in diesem Ordner.
"""
import os

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.stats import spearmanr

HERE = os.path.dirname(os.path.abspath(__file__))
LEVELS = [1, 2, 3, 4, 5]


def read_ratings(path: str, key: pd.DataFrame) -> pd.Series:
    ws = load_workbook(path, data_only=True)["Bewertung"]
    vals = {}
    for _, k in key.iterrows():
        v = ws[k["rating_cell"]].value
        vals[k["item_id"]] = int(v) if v not in (None, "") else np.nan
    return pd.Series(vals)


def weighted_kappa(a: np.ndarray, b: np.ndarray) -> float:
    n = len(LEVELS)
    idx = {v: i for i, v in enumerate(LEVELS)}
    obs = np.zeros((n, n))
    for x, y in zip(a, b):
        obs[idx[x], idx[y]] += 1
    obs /= obs.sum()
    exp = np.outer(obs.sum(axis=1), obs.sum(axis=0))
    w = np.array([[(i - j) ** 2 for j in range(n)] for i in range(n)]) / (n - 1) ** 2
    return 1 - (w * obs).sum() / (w * exp).sum()


def main():
    key = pd.read_csv(os.path.join(HERE, "key_DO_NOT_SHARE.csv"))
    ra = read_ratings(os.path.join(HERE, "Rater_A.xlsx"), key)
    rb = read_ratings(os.path.join(HERE, "Rater_B.xlsx"), key)
    df = key.set_index("item_id").assign(rater_a=ra, rater_b=rb)

    missing = df[df.rater_a.isna() | df.rater_b.isna()].index.tolist()
    if missing:
        print(f"[WARN] Fehlende Bewertungen: {missing}")
    both = df.dropna(subset=["rater_a", "rater_b"]).copy()
    both["human_mean"] = both[["rater_a", "rater_b"]].mean(axis=1)

    kappa = weighted_kappa(both.rater_a.astype(int).values, both.rater_b.astype(int).values)
    exact = (both.rater_a == both.rater_b).mean()
    within1 = ((both.rater_a - both.rater_b).abs() <= 1).mean()
    rho, p = spearmanr(both.human_mean, both.ragas_faithfulness)

    print(f"n = {len(both)} Eintraege mit zwei Bewertungen")
    print(f"Mensch-Mensch: gewichtetes kappa = {kappa:.3f} | exakte Uebereinstimmung = {exact:.0%} "
          f"| Abweichung <= 1 Stufe = {within1:.0%}")
    print(f"Mensch-Judge:  Spearman rho = {rho:.3f} (p = {p:.4f})")
    for src, g in both.groupby("source"):
        r, pp = spearmanr(g.human_mean, g.ragas_faithfulness)
        print(f"   {src:4s} n={len(g):2d}  rho = {r:.3f} (p = {pp:.3f})")
    print("\nMittel je Stratum (Mensch 1-5 vs. RAGAS 0-1):")
    print(both.groupby("stratum")[["human_mean", "ragas_faithfulness"]].mean().round(2).to_string())

    both.to_csv(os.path.join(HERE, "ratings_merged.csv"))
    print(f"\n-> {os.path.join(HERE, 'ratings_merged.csv')}")


if __name__ == "__main__":
    main()
